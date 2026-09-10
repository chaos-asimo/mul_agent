# -*- coding: utf-8 -*-
"""多用户龙虾Claw子项目 - PPT 生成服务（python-pptx，user_id 隔离）

流程：模型输出 [PPT_GEN]{...json...}[/PPT_GEN] → parse_ppt_json 解析 →
render_pptx 渲染（3 套内置主题，微软雅黑，封面+内容页版式，页数截断）→
保存到 files_dir(user_id) 并返回下载链接。
"""
import json
import logging
import os
import re
import subprocess
import sys
from typing import Any, Dict, List, Optional

from lobster_mu.db import BASE_DIR
from lobster_mu.paths import files_dir, gen_filename

logger = logging.getLogger(__name__)

MAX_SLIDES = 15
MIN_SLIDES = 5
MAX_BULLETS_PER_SLIDE = 6
MAX_BULLET_LEN = 60
MAX_CHART_CATEGORIES = 8
MAX_CHART_SERIES = 3

_LAYOUTS = {"bullets", "chart", "text_chart"}
_CHART_TYPES = {"bar", "line", "pie"}

_FONT = "微软雅黑"

# 内置主题：primary 主色 / accent 强调色 / bg 内容页底色 / text 正文色
THEMES = {
    "blue":   {"name": "商务蓝", "primary": (31, 78, 121),  "accent": (68, 114, 196),  "bg": (255, 255, 255), "text": (51, 51, 51)},
    "orange": {"name": "活力橙", "primary": (191, 79, 28),  "accent": (237, 125, 49),  "bg": (255, 255, 255), "text": (51, 51, 51)},
    "gray":   {"name": "简约灰", "primary": (68, 68, 68),   "accent": (128, 128, 128), "bg": (250, 250, 250), "text": (38, 38, 38)},
}

_THEME_ALIASES = {
    "blue": "blue", "商务蓝": "blue", "蓝": "blue",
    "orange": "orange", "活力橙": "orange", "橙": "orange",
    "gray": "gray", "grey": "gray", "简约灰": "gray", "灰": "gray",
}

_pptx_ready = False


def _ensure_pptx() -> bool:
    """确保 python-pptx 可用；缺失时安装到项目 packages/ 目录（清华镜像）"""
    global _pptx_ready
    if _pptx_ready:
        return True
    try:
        import pptx  # noqa: F401
        _pptx_ready = True
        return True
    except ImportError:
        pass
    pkg_dir = os.path.join(BASE_DIR, "packages")
    os.makedirs(pkg_dir, exist_ok=True)
    if pkg_dir not in sys.path:
        sys.path.insert(0, pkg_dir)
    try:
        import pptx  # noqa: F401
        _pptx_ready = True
        return True
    except ImportError:
        pass
    try:
        logger.warning("python-pptx 缺失，自动安装到 packages/ ...")
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "python-pptx",
             "--target", pkg_dir,
             "-i", "https://pypi.tuna.tsinghua.edu.cn/simple",
             "--trusted-host", "pypi.tuna.tsinghua.edu.cn"],
            timeout=600,
        )
        import pptx  # noqa: F401
        _pptx_ready = True
        return True
    except Exception as e:
        logger.error("python-pptx 自动安装失败: %s", e)
        return False


def _parse_chart(raw: Any) -> Optional[Dict[str, Any]]:
    """校验并清洗 chart 规格；非法返回 None。
    期望：{"type":"bar|line|pie","title":"...","categories":["A",...],"series":[{"name":"s1","values":[1,2]}]}"""
    if not isinstance(raw, dict):
        return None
    ctype = str(raw.get("type") or "").strip().lower()
    if ctype not in _CHART_TYPES:
        return None
    categories = raw.get("categories")
    if not isinstance(categories, list) or not categories:
        return None
    categories = [str(c).strip()[:20] for c in categories if str(c).strip()][:MAX_CHART_CATEGORIES]
    if not categories:
        return None
    series_raw = raw.get("series")
    if not isinstance(series_raw, list) or not series_raw:
        return None
    out_series: List[Dict[str, Any]] = []
    for s in series_raw[:MAX_CHART_SERIES]:
        if not isinstance(s, dict):
            continue
        values = s.get("values")
        if not isinstance(values, list):
            continue
        nums: List[float] = []
        for v in values[: len(categories)]:
            try:
                nums.append(float(v))
            except (TypeError, ValueError):
                nums.append(0.0)
        if not nums:
            continue
        # 与 categories 对齐长度（不足补 0）
        while len(nums) < len(categories):
            nums.append(0.0)
        out_series.append({"name": str(s.get("name") or "").strip()[:30], "values": nums})
    if not out_series:
        return None
    return {
        "type": ctype,
        "title": str(raw.get("title") or "").strip()[:60],
        "categories": categories,
        "series": out_series,
    }


def _repair_common_llm_json_errors(text: str) -> str:
    """修复 LLM 常见 JSON 书写错误：
    1. "layout":"bullets":[...] → "layout":"bullets","bullets":[...]（字段合并）
    2. 尾随逗号 ,] 或 ,}
    3. chart 对象后多写一个 } 导致 slide 提前闭合：]}]}},"note":... → }]},"note":...
    """
    # 字段合并：layout 的值后面紧跟冒号+数组/字符串，说明 bullets 键被吞了
    text = re.sub(
        r'"layout"\s*:\s*"(bullets|chart|text_chart)"\s*:\s*(\[)',
        r'"layout":"\1","bullets":\2',
        text,
    )
    # slide 提前闭合：series 数组 + chart 对象闭合后又多了一个 } 或 ]，导致 note/bullets 等字段悬空
    prev = None
    while prev != text:  # 可能连续多处，循环修复
        prev = text
        text = re.sub(r'(\]\s*\})\s*[}\]]\s*,\s*"(note|bullets|layout|chart|heading)"\s*:', r'\1,"\2":', text)
    # 尾随逗号
    text = re.sub(r",\s*([}\]])", r"\1", text)
    return text


def _close_truncated_json(text: str) -> str:
    """为被截断的 JSON 补齐未闭合的字符串/数组/对象括号"""
    stack: List[str] = []
    in_str = False
    esc = False
    for ch in text:
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
        else:
            if ch == '"':
                in_str = True
            elif ch in "{[":
                stack.append(ch)
            elif ch == "}":
                if stack and stack[-1] == "{":
                    stack.pop()
            elif ch == "]":
                if stack and stack[-1] == "[":
                    stack.pop()
    out = text.rstrip().rstrip(",")
    if in_str:
        out += '"'
    out += "".join("}" if c == "{" else "]" for c in reversed(stack))
    return out


def _try_loads(text: str) -> Optional[Any]:
    try:
        return json.loads(text)
    except Exception:
        return None


def parse_ppt_json(payload: str) -> Optional[Dict[str, Any]]:
    """解析 [PPT_GEN] 标记内的 JSON 大纲；非法返回 None。
    容错：截取最外层 {} 区间、修正常见尾随逗号、layout/bullets 字段合并。"""
    if not payload:
        return None
    text = payload.strip()
    start = text.find("{")
    end = text.rfind("}")
    if start == -1:
        return None
    if end == -1 or end <= start:
        end = len(text)  # 无任何闭合括号：整体视为被截断，进入挽救流程
    text = text[start:end + 1] if end < len(text) else text[start:]
    data = _try_loads(text)
    if data is None:
        repaired = _repair_common_llm_json_errors(text)
        data = _try_loads(repaired)
        if data is None:
            # 截断挽救：先整体闭合；失败则回退到最近一个完整 slide 对象再闭合
            data = _try_loads(_close_truncated_json(repaired))
            if data is None:
                # 记录完整 payload 便于排查语法错误
                try:
                    import time as _t
                    dbg = os.path.join(BASE_DIR, "data", f"ppt_parse_fail_{int(_t.time())}.json")
                    os.makedirs(os.path.dirname(dbg), exist_ok=True)
                    with open(dbg, "w", encoding="utf-8") as f:
                        f.write(text)  # 保存原始 payload（修复前），便于定位模型真实错误
                    logger.warning("PPT 大纲解析失败，payload 已保存: %s", dbg)
                except Exception:
                    pass
                for m in reversed(list(re.finditer(r"\}\s*,", repaired))):
                    data = _try_loads(_close_truncated_json(repaired[: m.start() + 1]))
                    if data is not None:
                        logger.warning("PPT 大纲被截断，已回退保留完整页 | 原始尾部: %s", repaired[-200:])
                        break
            if data is None:
                return None
    if not isinstance(data, dict):
        return None
    slides = data.get("slides")
    if not isinstance(slides, list) or not slides:
        return None
    title = str(data.get("title") or "演示文稿").strip()[:60] or "演示文稿"
    out_slides: List[Dict[str, Any]] = []
    for s in slides:
        if not isinstance(s, dict):
            continue
        heading = str(s.get("heading") or "").strip()[:60]
        bullets = s.get("bullets")
        if not isinstance(bullets, list):
            bullets = [str(bullets)] if bullets else []
        bullets = [str(b).strip()[:MAX_BULLET_LEN] for b in bullets if str(b).strip()]
        bullets = bullets[:MAX_BULLETS_PER_SLIDE]
        chart = _parse_chart(s.get("chart"))
        layout = str(s.get("layout") or "").strip().lower()
        if layout not in _LAYOUTS:
            layout = "text_chart" if chart and bullets else ("chart" if chart else "bullets")
        if layout in ("chart", "text_chart") and not chart:
            layout = "bullets"
        if not heading and not bullets and not chart:
            continue
        out_slides.append({
            "heading": heading or "（无标题）",
            "bullets": bullets,
            "note": str(s.get("note") or "").strip()[:500],
            "layout": layout,
            "chart": chart,
        })
    if not out_slides:
        return None
    truncated = len(out_slides) > MAX_SLIDES
    if truncated:
        out_slides = out_slides[:MAX_SLIDES]
    theme_key = _THEME_ALIASES.get(str(data.get("theme") or "").strip().lower(), None) \
        or _THEME_ALIASES.get(str(data.get("theme") or "").strip(), "blue")
    return {
        "title": title,
        "subtitle": str(data.get("subtitle") or "").strip()[:80],
        "theme": theme_key,
        "slides": out_slides,
        "truncated": truncated,
    }


def _set_run_font(run, size_pt: int, color, bold: bool = False):
    """设置中英文字体（微软雅黑）与字号颜色"""
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn

    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = _FONT
    run.font.color.rgb = RGBColor(*color)
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", _FONT)


def _render_bullets(tf, bullets, theme, size: int = 20):
    """在文本框内渲染要点列表（强调色圆点 + 正文）"""
    from pptx.util import Pt

    tf.word_wrap = True
    if not bullets:
        run = tf.paragraphs[0].add_run()
        run.text = " "
        _set_run_font(run, size, tuple(theme["text"]))
        return
    for i, b in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(12)
        para.line_spacing = 1.25
        dot = para.add_run()
        dot.text = "• "
        _set_run_font(dot, size, tuple(theme["accent"]), bold=True)
        run = para.add_run()
        run.text = b
        _set_run_font(run, size, tuple(theme["text"]))


def _style_chart_text(chart, size: int = 11):
    """统一图表整体字体（图例/坐标轴/数据标签继承）"""
    from pptx.util import Pt

    chart.font.size = Pt(size)
    chart.font.name = _FONT


def _theme_palette(theme: Dict[str, Any]) -> List[tuple]:
    """由主题色生成 8 色调色板（主/强调色的明暗变体），供图表分类着色"""
    def lighten(c, f):
        return tuple(int(x + (255 - x) * f) for x in c)

    def darken(c, f):
        return tuple(int(x * (1 - f)) for x in c)

    a, p = tuple(theme["accent"]), tuple(theme["primary"])
    return [a, p, lighten(a, 0.45), lighten(p, 0.35), darken(a, 0.25), lighten(p, 0.6), darken(p, 0.25), (170, 170, 170)]


def _add_chart(slide, chart_spec: Dict[str, Any], theme, x, y, cx, cy):
    """添加原生图表（bar/line/pie），返回 GraphicFrame"""
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    chart_data = CategoryChartData()
    chart_data.categories = chart_spec["categories"]
    for s in chart_spec["series"]:
        chart_data.add_series(s["name"] or "系列", s["values"])

    type_map = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }
    frame = slide.shapes.add_chart(type_map[chart_spec["type"]], x, y, cx, cy, chart_data)
    chart = frame.chart
    _style_chart_text(chart, 11)

    if chart_spec.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = chart_spec["title"]
        for p in chart.chart_title.text_frame.paragraphs:
            for r in p.runs:
                _set_run_font(r, 14, tuple(theme["primary"]), bold=True)
    else:
        chart.has_title = False

    if chart_spec["type"] == "pie":
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        # 饼图显示数据标签
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.show_value = True
        plot.data_labels.font.size = None
    else:
        chart.has_legend = len(chart_spec["series"]) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False

    # 系列配色：主题 8 色调色板轮换
    palette = _theme_palette(theme)
    if chart_spec["type"] == "pie":
        for i, pt in enumerate(chart.plots[0].series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor(*palette[i % len(palette)])
    elif chart_spec["type"] == "line":
        # 折线图：设置线条颜色与数据点标记（fill 对线条无效）
        from pptx.enum.chart import XL_MARKER_STYLE
        from pptx.util import Pt as _Pt
        for i, series in enumerate(chart.series):
            color = RGBColor(*palette[i % len(palette)])
            series.format.line.color.rgb = color
            series.format.line.width = _Pt(2.25)
            try:
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = color
                series.marker.format.line.color.rgb = color
            except Exception:
                pass
    else:
        for i, series in enumerate(chart.series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor(*palette[i % len(palette)])
    return frame


def _add_footer(slide, idx: int, total: int, deck_title: str, theme, sw, sh):
    """页脚：左下标题 + 右下页码"""
    from pptx.util import Inches
    from pptx.enum.text import PP_ALIGN

    fbox = slide.shapes.add_textbox(Inches(0.5), sh - Inches(0.45), sw - Inches(1.0), Inches(0.3))
    tf = fbox.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = deck_title[:30]
    _set_run_font(r, 10, (150, 150, 150))
    # 页码（独立文本框右对齐）
    pbox = slide.shapes.add_textbox(sw - Inches(1.2), sh - Inches(0.45), Inches(0.9), Inches(0.3))
    pp = pbox.text_frame.paragraphs[0]
    pp.alignment = PP_ALIGN.RIGHT
    pr = pp.add_run()
    pr.text = f"{idx} / {total}"
    _set_run_font(pr, 10, (150, 150, 150))


def render_pptx(data: Dict[str, Any], out_path: str) -> int:
    """渲染 PPTX 到 out_path，返回实际页数（含封面）。data 需为 parse_ppt_json 的产物。"""
    if not _ensure_pptx():
        raise RuntimeError("python-pptx 未安装且自动安装失败")
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    theme = THEMES.get(data.get("theme") or "blue", THEMES["blue"])
    primary = RGBColor(*theme["primary"])
    accent = RGBColor(*theme["accent"])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    sw, sh = prs.slide_width, prs.slide_height

    # ---- 封面页：主色满铺 + 强调色横线 + 居中标题/副标题 ----
    cover = prs.slides.add_slide(blank)
    band = cover.shapes.add_shape(1, 0, 0, sw, sh)  # 1 = 矩形
    band.fill.solid()
    band.fill.fore_color.rgb = primary
    band.line.fill.background()
    # 强调色装饰横线
    line = cover.shapes.add_shape(1, int(sw / 2) - Inches(1), Inches(4.15), Inches(2), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()
    box = cover.shapes.add_textbox(Inches(1), Inches(2.6), sw - Inches(2), Inches(1.6))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run_font(p.add_run(), 40, (255, 255, 255), bold=True)
    p.runs[0].text = data["title"]
    if data.get("subtitle"):
        sbox = cover.shapes.add_textbox(Inches(1), Inches(4.4), sw - Inches(2), Inches(1.0))
        sp = sbox.text_frame.paragraphs[0]
        sp.alignment = PP_ALIGN.CENTER
        r = sp.add_run()
        r.text = data["subtitle"]
        _set_run_font(r, 20, (230, 230, 230))

    total = 1 + len(data["slides"])

    # ---- 内容页 ----
    for idx, s in enumerate(data["slides"], start=2):
        slide = prs.slides.add_slide(blank)
        # 顶部标题条
        bar = slide.shapes.add_shape(1, 0, 0, sw, Inches(1.0))
        bar.fill.solid()
        bar.fill.fore_color.rgb = primary
        bar.line.fill.background()
        # 标题条下强调色细线
        underline = slide.shapes.add_shape(1, 0, Inches(1.0), sw, Inches(0.05))
        underline.fill.solid()
        underline.fill.fore_color.rgb = accent
        underline.line.fill.background()
        tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), sw - Inches(1.2), Inches(0.7))
        tp = tbox.text_frame.paragraphs[0]
        r = tp.add_run()
        r.text = s["heading"]
        _set_run_font(r, 26, (255, 255, 255), bold=True)

        layout = s.get("layout") or "bullets"
        body_top = Inches(1.35)
        body_h = sh - Inches(2.1)

        if layout == "chart" and s.get("chart"):
            # 整页图表（可选要点在底部）
            if s["bullets"]:
                _add_chart(slide, s["chart"], theme, Inches(1.0), body_top, sw - Inches(2.0), Inches(3.6))
                body = slide.shapes.add_textbox(Inches(1.2), Inches(5.15), sw - Inches(2.4), Inches(1.3))
                _render_bullets(body.text_frame, s["bullets"][:3], theme, size=14)
            else:
                _add_chart(slide, s["chart"], theme, Inches(1.0), body_top, sw - Inches(2.0), body_h)
        elif layout == "text_chart" and s.get("chart"):
            # 左文右图混排
            left = slide.shapes.add_textbox(Inches(0.6), body_top, Inches(5.4), body_h)
            left.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            _render_bullets(left.text_frame, s["bullets"], theme, size=16)
            _add_chart(slide, s["chart"], theme, Inches(6.3), body_top, Inches(6.5), body_h)
        else:
            # 纯要点（垂直居中）
            body = slide.shapes.add_textbox(Inches(0.8), body_top, sw - Inches(1.6), body_h)
            body.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            _render_bullets(body.text_frame, s["bullets"], theme, size=20)

        _add_footer(slide, idx, total, data["title"], theme, sw, sh)

        if s.get("note"):
            slide.notes_slide.notes_text_frame.text = s["note"]

    prs.save(out_path)
    return 1 + len(data["slides"])


def generate_ppt(user_id: int, data: Dict[str, Any]) -> Dict[str, Any]:
    """完整生成：渲染 + 保存到用户产物目录 + 返回下载信息"""
    if not _ensure_pptx():
        return {"success": False, "error": "python-pptx 未安装且自动安装失败"}
    try:
        filename = gen_filename("presentation.pptx")
        out_path = os.path.join(files_dir(user_id), filename)
        slide_count = render_pptx(data, out_path)
        return {
            "success": True,
            "title": data["title"],
            "filename": filename,
            "slide_count": slide_count,
            "truncated": bool(data.get("truncated")),
            "theme_name": THEMES.get(data.get("theme") or "blue", THEMES["blue"])["name"],
            "download_url": f"/api/lobster-mu/script/files/download/{filename}",
        }
    except Exception as e:
        logger.error("PPT 渲染失败: %s", e)
        return {"success": False, "error": f"渲染失败: {str(e)[:120]}"}
