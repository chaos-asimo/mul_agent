# -*- coding: utf-8 -*-
"""多用户龙虾Claw子项目 - 文件上传与文本提取（按 user_id 隔离存储）"""
import os
import json
from datetime import datetime

from lobster_mu.paths import uploads_dir, gen_filename

# 10MB 上传限制
MAX_UPLOAD_SIZE = 10 * 1024 * 1024


def _now() -> str:
    return datetime.now().isoformat(sep=" ", timespec="seconds")


def extract_text(file_path: str, ext: str, raw: bytes = None) -> str:
    """按扩展名提取文本内容（复制自原版逻辑）"""
    if ext in ['.txt', '.md', '.json', '.py', '.js', '.html', '.css', '.xml', '.csv']:
        if raw is not None:
            return raw.decode('utf-8', errors='replace')
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            return f.read()
    if ext in ['.doc', '.docx']:
        try:
            from docx import Document
            doc = Document(file_path)
            return '\n'.join([para.text for para in doc.paragraphs])
        except ImportError:
            return "[需要安装 python-docx 库来解析Word文档]"
        except Exception as e:
            return f"[解析Word文档失败: {str(e)}]"
    if ext == '.pdf':
        try:
            import fitz
            doc = fitz.open(file_path)
            text = '\n'.join([page.get_text() for page in doc])
            if len(text.strip()) < 100:
                try:
                    import pytesseract
                    from PIL import Image
                    ocr_text = ""
                    for page in doc:
                        pix = page.get_pixmap()
                        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        ocr_text += pytesseract.image_to_string(img, lang='chi_sim') + '\n'
                    text = ocr_text
                except Exception:
                    pass
            return text
        except ImportError:
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(file_path)
                return '\n'.join([page.extract_text() or '' for page in reader.pages])
            except ImportError:
                return "[需要安装 PyMuPDF(fitz) 或 PyPDF2 库来解析PDF文档]"
            except Exception as e:
                return f"[使用PyPDF2解析PDF失败: {str(e)}]"
        except Exception as e:
            return f"[解析PDF文档失败: {str(e)}]"
    if ext in ['.xls', '.xlsx']:
        try:
            import pandas as pd
            df = pd.read_excel(file_path)
            return df.to_string()
        except ImportError:
            return "[需要安装 pandas 和 openpyxl 库来解析Excel文件]"
        except Exception as e:
            return f"[解析Excel文件失败: {str(e)}]"
    if ext in ['.ppt', '.pptx']:
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            text = ''
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        text += shape.text + '\n'
            return text
        except ImportError:
            return "[需要安装 python-pptx 库来解析PowerPoint文件]"
        except Exception as e:
            return f"[解析PowerPoint文件失败: {str(e)}]"
    return f"[不支持的文件类型 {ext}]"


def save_upload(user_id: int, filename: str, content: bytes) -> dict:
    """保存上传文件到用户目录并提取文本，返回文件信息"""
    ext = os.path.splitext(filename or "")[1][:16]
    target_dir = uploads_dir(user_id)
    unique_name = gen_filename(filename)
    file_path = os.path.join(target_dir, unique_name)
    with open(file_path, 'wb') as f:
        f.write(content)
    text_content = extract_text(file_path, ext.lower(), raw=content)
    return {
        "filename": filename,
        "path": file_path,
        "size": len(content),
        "content": text_content,
        "timestamp": _now(),
    }


def save_upload_text(user_id: int, filename: str, content: bytes) -> dict:
    """上传并返回解析文本（v2 前端 /upload/text 使用）"""
    ext = os.path.splitext(filename or "")[1][:16].lower()
    target_dir = uploads_dir(user_id)
    unique_name = gen_filename(filename)
    file_path = os.path.join(target_dir, unique_name)
    with open(file_path, 'wb') as f:
        f.write(content)
    print(f"[DEBUG] mu上传文件: {filename}, 大小: {len(content)} bytes, 扩展名: {ext}")
    print(f"[DEBUG] mu文件已保存到: {file_path}")
    text_content = extract_text(file_path, ext, raw=content)
    print(f"[DEBUG] mu解析内容长度: {len(text_content)}")
    return {"success": True, "filename": filename, "content": text_content}
