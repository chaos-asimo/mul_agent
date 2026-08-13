import os, json

from .vector_store import VectorStore
from .text_chunker import TextChunker


class KnowledgeBase:
    """RAG knowledge base manager.

    Coordinates file parsing, text chunking, embedding generation and vector
    storage to provide a high-level interface for adding documents and
    performing semantic search.
    """

    def __init__(self, model_manager=None):
        """Initialize the knowledge base.

        Args:
            model_manager: Optional ``ModelManager`` instance used to look up
                configured embedding models. May be set later via
                :meth:`set_model_manager`.
        """
        self.vector_store = VectorStore()
        self.text_chunker = TextChunker()
        self.model_manager = model_manager

    def set_model_manager(self, model_manager):
        """Set or update the model manager reference.

        Args:
            model_manager: ``ModelManager`` instance used to look up
                configured embedding models.
        """
        self.model_manager = model_manager
        print(f"[KnowledgeBase] model_manager 已设置")

    def _get_embedding_adapter(self):
        """Return an embedding adapter for the first enabled embedding model.

        Looks up enabled models from ``model_manager.get_all()`` whose
        ``model_type`` is ``"embedding"`` and whose ``api_key`` is non-empty,
        then builds an adapter via ``create_embedding_adapter``.

        Returns:
            An LLMAdapter with an ``embeddings(texts)`` method, or ``None`` if
            no embedding model is configured or no model manager is set.
        """
        if self.model_manager is None:
            print(f"[KnowledgeBase] model_manager 未配置")
            return None
        try:
            from engine.agent_worker import create_embedding_adapter
        except ImportError as e:
            print(f"[KnowledgeBase] 无法导入 create_embedding_adapter: {e}")
            return None

        try:
            models = self.model_manager.get_all()
        except Exception as e:
            print(f"[KnowledgeBase] 获取模型列表失败: {e}")
            return None

        for model_config in models:
            if not getattr(model_config, "enabled", False):
                continue
            if getattr(model_config, "model_type", "") != "embedding":
                continue
            if not getattr(model_config, "api_key", ""):
                continue
            try:
                adapter = create_embedding_adapter(model_config)
                if adapter is not None:
                    print(f"[KnowledgeBase] 使用 embedding 模型: {model_config.model_name}")
                    return adapter
            except Exception as e:
                print(f"[KnowledgeBase] 创建 embedding adapter 失败: {e}")
                continue
        print(f"[KnowledgeBase] 未找到可用的 embedding 模型")
        return None

    def parse_file(self, file_path, file_ext):
        """Parse a file and return its extracted text content.

        Supports: .txt, .md, .json, .py, .js, .html, .css, .xml, .csv
        (utf-8 decode), .doc/.docx (python-docx), .pdf (PyMuPDF/fitz),
        .xls/.xlsx (pandas) and .ppt/.pptx (python-pptx).

        Args:
            file_path: Absolute path to the file on disk.
            file_ext: File extension including the leading dot (case-insensitive).

        Returns:
            Extracted text as a string. If parsing fails or the file type is
            unsupported, a descriptive error string wrapped in brackets is
            returned.
        """
        ext = (file_ext or "").lower()
        try:
            if ext in ['.txt', '.md', '.json', '.py', '.js', '.html', '.css', '.xml', '.csv']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    text_content = f.read()
            elif ext in ['.doc', '.docx']:
                try:
                    from docx import Document
                    doc = Document(file_path)
                    text_content = '\n'.join([para.text for para in doc.paragraphs])
                except ImportError:
                    text_content = f"[需要安装 python-docx 库来解析Word文档]"
                except Exception as e:
                    text_content = f"[解析Word文档失败: {str(e)}]"
            elif ext == '.pdf':
                try:
                    import fitz
                    doc = fitz.open(file_path)
                    text_content = '\n'.join([page.get_text() for page in doc])
                except ImportError:
                    text_content = f"[需要安装 PyMuPDF(fitz) 库来解析PDF文档]"
                except Exception as e:
                    text_content = f"[解析PDF文档失败: {str(e)}]"
            elif ext in ['.xls', '.xlsx']:
                try:
                    import pandas as pd
                    df = pd.read_excel(file_path)
                    text_content = df.to_string()
                except ImportError:
                    text_content = f"[需要安装 pandas 和 openpyxl 库来解析Excel文件]"
                except Exception as e:
                    text_content = f"[解析Excel文件失败: {str(e)}]"
            elif ext in ['.ppt', '.pptx']:
                try:
                    from pptx import Presentation
                    prs = Presentation(file_path)
                    text_content = ''
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, 'text'):
                                text_content += shape.text + '\n'
                except ImportError:
                    text_content = f"[需要安装 python-pptx 库来解析PowerPoint文件]"
                except Exception as e:
                    text_content = f"[解析PowerPoint文件失败: {str(e)}]"
            else:
                text_content = f"[不支持的文件类型 {file_ext}]"
        except Exception as e:
            text_content = f"[提取文件内容失败: {str(e)}]"
        return text_content

    def add_document(self, file_path, filename, file_size=None):
        """Parse, chunk, embed and store a file in the knowledge base.

        Args:
            file_path: Absolute path to the file on disk.
            filename: Original file name to store as metadata.
            file_size: Optional size of the file in bytes. When ``None`` the
                size is read from disk via ``os.path.getsize``.

        Returns:
            Dict with ``doc_id``, ``filename``, ``chunk_count`` and
            ``file_size``.

        Raises:
            ValueError: If no embedding model is configured.
        """
        doc_id = os.urandom(16).hex()
        if file_size is None:
            try:
                file_size = os.path.getsize(file_path)
            except Exception:
                file_size = 0

        file_ext = os.path.splitext(filename or file_path)[1]
        print(f"[KnowledgeBase] 开始处理文件: {filename} ({file_ext})")
        text_content = self.parse_file(file_path, file_ext)
        return self._ingest_text(doc_id, filename, file_size, text_content)

    def add_text_content(self, text, filename="手动添加", source="manual"):
        """Chunk, embed and store raw text in the knowledge base.

        Args:
            text: Raw text content to add.
            filename: Display name for the entry. Defaults to ``"手动添加"``.
            source: Source label stored in metadata. Defaults to ``"manual"``.

        Returns:
            Dict with ``doc_id``, ``filename``, ``chunk_count`` and
            ``file_size``.

        Raises:
            ValueError: If no embedding model is configured.
        """
        doc_id = os.urandom(16).hex()
        file_size = len(text.encode('utf-8')) if text else 0
        print(f"[KnowledgeBase] 添加文本内容: {filename} (source={source})")
        return self._ingest_text(doc_id, filename, file_size, text)

    def _ingest_text(self, doc_id, filename, file_size, text_content):
        """Internal helper: chunk, embed and store already-extracted text.

        Args:
            doc_id: Pre-generated document id.
            filename: Display name for the entry.
            file_size: Size value to store as metadata.
            text_content: Extracted text to ingest.

        Returns:
            Dict with ``doc_id``, ``filename``, ``chunk_count`` and
            ``file_size``.

        Raises:
            ValueError: If no embedding model is configured.
        """
        chunks = self.text_chunker.chunk(text_content or "")
        print(f"[KnowledgeBase] 文件 {filename} 分块完成，共 {len(chunks)} 块")

        if not chunks:
            print(f"[KnowledgeBase] 文件 {filename} 无有效内容")
            self.vector_store.add_document(doc_id, filename, file_size, [])
            return {
                'doc_id': doc_id,
                'filename': filename,
                'chunk_count': 0,
                'file_size': file_size,
            }

        adapter = self._get_embedding_adapter()
        if adapter is None:
            raise ValueError("未配置embedding模型，请先在模型配置中添加embedding类型的模型")

        try:
            embeddings = adapter.embeddings([chunk['content'] for chunk in chunks])
        except Exception as e:
            print(f"[KnowledgeBase] 生成 embedding 失败: {e}")
            raise

        chunk_data = []
        for chunk, embedding in zip(chunks, embeddings):
            chunk_data.append({
                'chunk_id': os.urandom(16).hex(),
                'chunk_index': chunk['chunk_index'],
                'content': chunk['content'],
                'embedding': embedding,
                'metadata': {
                    'start_pos': chunk.get('start_pos', 0),
                    'end_pos': chunk.get('end_pos', 0),
                },
            })

        self.vector_store.add_document(doc_id, filename, file_size, chunk_data)
        print(f"[KnowledgeBase] 文件 {filename} 已入库 (doc_id={doc_id})")
        return {
            'doc_id': doc_id,
            'filename': filename,
            'chunk_count': len(chunk_data),
            'file_size': file_size,
        }

    def search(self, query, top_k=5, threshold=0.3):
        """Semantic search against the knowledge base.

        Args:
            query: Query text to embed and search with.
            top_k: Maximum number of results to return.
            threshold: Minimum cosine similarity score for a result to be kept.

        Returns:
            List of result dicts from ``VectorStore.search``, or an empty list
            if no embedding model is configured or an error occurs.
        """
        from .rag_logger import (
            rag_logger, log_retrieval_start, log_retrieval_result,
            log_error, log_debug, log_embedding_status
        )
        import time
        
        if not query:
            rag_logger.warning("[KnowledgeBase] 检索查询为空，返回 0 条结果")
            return []
        
        retrieval_start = log_retrieval_start(query, top_k, threshold, "KnowledgeBase")
        
        adapter = self._get_embedding_adapter()
        if adapter is None:
            rag_logger.warning("[KnowledgeBase] 检索中止：未配置 embedding 模型")
            return []
        
        # Stage 1: Generate query embedding
        embed_start = time.time()
        try:
            query_embedding = adapter.embed_query(query)
            embed_ms = (time.time() - embed_start) * 1000
            log_embedding_status(query, len(query_embedding), embed_ms, "KnowledgeBase")
        except Exception as e:
            log_error("KnowledgeBase", "embed_query", e, f"query='{query[:50]}'")
            return []
        
        # Stage 2: Vector similarity search
        search_start = time.time()
        try:
            results = self.vector_store.search(query_embedding, top_k, threshold)
            search_ms = (time.time() - search_start) * 1000
            
            total_elapsed = (time.time() - retrieval_start) * 1000
            rag_logger.info(
                f"[KnowledgeBase] 检索完成 | "
                f"向量化={embed_ms:.1f}ms | 相似度检索={search_ms:.1f}ms | "
                f"总计={total_elapsed:.1f}ms | 命中={len(results)}条"
            )
            
            if not results:
                rag_logger.info(
                    f"[KnowledgeBase] 无检索结果 | "
                    f"查询='{query[:50]}' | threshold={threshold} | "
                    f"可能原因：知识库为空、查询无相关内容、或阈值过高"
                )
            else:
                rag_logger.debug(
                    f"[KnowledgeBase] 结果摘要 | "
                    f"最高得分={results[0]['score']:.4f} | "
                    f"最低得分={results[-1]['score']:.4f} | "
                    f"文档={set(r.get('filename', '?') for r in results)}"
                )
            
            return results
        except Exception as e:
            search_ms = (time.time() - search_start) * 1000
            log_error("KnowledgeBase", "vector_search", e, 
                f"query='{query[:50]}', embed_ms={embed_ms:.1f}, search_ms={search_ms:.1f}"
            )
            return []

    def delete_document(self, doc_id):
        """Delete a document and all of its chunks.

        Args:
            doc_id: Unique identifier of the document to delete.

        Returns:
            The number of deleted chunks.
        """
        deleted = self.vector_store.delete_document(doc_id)
        print(f"[KnowledgeBase] 删除文档 doc_id={doc_id}，删除块数: {deleted}")
        return deleted

    def list_documents(self):
        """Return all documents ordered by upload time (newest first).

        Returns:
            List of document dicts from ``VectorStore.list_documents``.
        """
        return self.vector_store.list_documents()

    def get_stats(self):
        """Return aggregate statistics about the knowledge base.

        Returns:
            Dict with ``total_documents``, ``total_chunks`` and
            ``db_size_bytes``.
        """
        return self.vector_store.get_stats()
